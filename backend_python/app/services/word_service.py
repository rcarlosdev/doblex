import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.models.ot import Ot

def generate_ot_word(ot: Ot) -> io.BytesIO:
    """
    Genera un informe técnico de Orden de Trabajo en formato Word (.docx).
    """
    doc = Document()

    # Título principal
    title = doc.add_heading(f"INFORME TÉCNICO - {ot.codigo}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_sub = doc.add_paragraph("DOBLEX S.A.S. - Sistema de Gestión de Obra Civil (SMU)")
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.runs[0].font.size = Pt(10)
    p_sub.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # 1. Datos Generales
    doc.add_heading("1. Información General de la Orden", level=1)
    table = doc.add_table(rows=6, cols=2)
    table.style = 'Table Grid'

    assigned_name = ot.assigned_user.name if ot.assigned_user else "Sin Asignar"
    cuadrilla_name = ot.cuadrilla.nombre if ot.cuadrilla else "N/A"

    info_rows = [
        ("Código de OT:", ot.codigo),
        ("Descripción:", ot.descripcion),
        ("Ubicación / Sitio:", f"{ot.sitio or ''} ({ot.ubicacion})"),
        ("Técnico Asignado:", f"{assigned_name} (Cuadrilla: {cuadrilla_name})"),
        ("Estado / Progreso:", f"{ot.estado.upper()} - {ot.progreso}%"),
        ("Prioridad / SLA:", f"{ot.prioridad} ({ot.tipo_ubicacion}) - Límite: {ot.fecha_limite_sla or 'N/A'}")
    ]

    for i, (k, v) in enumerate(info_rows):
        row = table.rows[i]
        row.cells[0].paragraphs[0].add_run(k).bold = True
        row.cells[1].paragraphs[0].add_run(str(v))

    doc.add_paragraph()

    # 2. Cierre y Causa Falla
    if ot.estado in ["solucionada", "finalizada"]:
        doc.add_heading("2. Cierre Técnico y Diagnóstico", level=1)
        doc.add_paragraph(f"Causa de la Falla: {ot.causa_falla or 'No especificada'}")
        if ot.observaciones_cierre:
            doc.add_paragraph(f"Observaciones de Cierre: {ot.observaciones_cierre}")
        doc.add_paragraph()

    # 3. Repuestos e Insumos
    if ot.repuestos:
        doc.add_heading("3. Insumos y Repuestos Utilizados", level=1)
        rep_table = doc.add_table(rows=len(ot.repuestos) + 1, cols=3)
        rep_table.style = 'Table Grid'

        # Cabecera
        h_row = rep_table.rows[0]
        h_row.cells[0].paragraphs[0].add_run("Ítem / Repuesto").bold = True
        h_row.cells[1].paragraphs[0].add_run("Cantidad").bold = True
        h_row.cells[2].paragraphs[0].add_run("Unidad").bold = True

        for idx, rep in enumerate(ot.repuestos, start=1):
            r = rep_table.rows[idx]
            r.cells[0].paragraphs[0].add_run(rep.nombre_item)
            r.cells[1].paragraphs[0].add_run(str(rep.cantidad))
            r.cells[2].paragraphs[0].add_run(rep.unidad_medida)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
