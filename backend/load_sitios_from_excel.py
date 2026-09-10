"""
Script ETL para cargar y sincronizar la base de datos de Sitios (Estaciones Base)
a partir de la primera pestaña ('Base de dato transporte') de docs/ARCHIVO FACTURACIÓN.xlsx
"""
import os
import sys
from datetime import datetime
import openpyxl

# Configurar el path del backend
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(BASE_DIR, "backend"))

from app.db.base import Base
from app.db.session import engine, SessionLocal
from app.models.sitio import Sitio
from app.models.ot import Ot

EXCEL_PATH = os.path.join(BASE_DIR, "docs", "ARCHIVO FACTURACIÓN.xlsx")

def parse_float_safe(val):
    if val is None:
        return None, None
    if isinstance(val, (int, float)):
        return float(val), str(val)
    if isinstance(val, datetime):
        # Cuando Excel convierte rangos como '2-5' en '2026-05-02'
        text_rep = f"{val.day}-{val.month}"
        return None, text_rep
    val_str = str(val).strip().replace(",", ".")
    if not val_str or val_str.lower() in ["#n/a", "#n/d", "nan", "null", "none", "0", "-", "--"]:
        if val_str == "0":
            return 0.0, "0"
        return None, None
    try:
        return float(val_str), str(val).strip()
    except (ValueError, TypeError):
        return None, str(val).strip()

def clean_str(val):
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    s_lower = s.lower()
    # Descartar errores de fórmula de Excel, ceros espurios o marcadores de no-aplica
    if s_lower in ["#n/a", "#n/d", "#ref!", "#valor!", "#value!", "0", "-", "--", "none", "null", "n/a", "no aplica", "nan", "#div/0!", "#num!", "#name?"]:
        return None
    if "#n/a" in s_lower or "#n/d" in s_lower or "#ref!" in s_lower:
        return None
    return s

def clean_email(val):
    s = clean_str(val)
    if not s:
        return None
    s = s.lower().strip()
    if "@" not in s or "#" in s or " " in s:
        return None
    return s

def load_sitios():
    if not os.path.exists(EXCEL_PATH):
        print(f"[ERROR] No se encontró el archivo Excel en: {EXCEL_PATH}")
        sys.exit(1)

    print("==================================================================")
    print(" [ETL] INICIANDO CARGA DE SITIOS DESDE EXCEL DE FACTURACIÓN")
    print(f" Archivo origen: {EXCEL_PATH}")
    print("==================================================================")

    # 1. Asegurar que las tablas existan en la BD
    print("\n[Paso 1/4] Creando / verificando tablas en la base de datos...")
    Base.metadata.create_all(bind=engine)

    # 2. Cargar Excel
    print("\n[Paso 2/4] Abriendo archivo Excel y leyendo pestaña 'Base de dato transporte'...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet_name = wb.sheetnames[0]  # Primer pestaña: 'Base de dato transporte'
    ws = wb[sheet_name]
    print(f"  -> Pestaña cargada: '{sheet_name}' (Total filas: {ws.max_row})")

    # Mapeo de columnas de la fila 19
    # Col 1: ZONA, 2: sitio, 3: Código Transporte LPU, 4: Km, 5: CIUDAD BASE,
    # 6: TRANSPORTE ESPECIAL, 7: ZONA TECNICA, 8: New_SO, 9: jefe_zona, 10: Ing Soporte,
    # 11: Estructura, 12: Altura Estructura, 13: FACTURADORA, 14: MUNICIPIO, 15: UBICACIÓN,
    # 16: CORREOS SO, 17: CORREO JEFE DE ZONA, 18: CORREO Ing.SOPORTE

    db = SessionLocal()
    inserted_count = 0
    updated_count = 0
    skipped_count = 0

    try:
        print("\n[Paso 3/4] Procesando filas e insertando/actualizando registros...")
        # Cache de sitios existentes para optimizar velocidad
        existing_sitios = {s.nombre: s for s in db.query(Sitio).all()}

        batch_size = 200
        current_batch = 0

        for r in range(20, ws.max_row + 1):
            sitio_raw = ws.cell(r, 2).value
            if not sitio_raw or not str(sitio_raw).strip():
                skipped_count += 1
                continue

            nombre = str(sitio_raw).strip().upper()
            zona = clean_str(ws.cell(r, 1).value)
            if zona:
                zona = zona.upper()

            cod_transp = clean_str(ws.cell(r, 3).value)
            km_num, km_txt = parse_float_safe(ws.cell(r, 4).value)
            ciudad_base = clean_str(ws.cell(r, 5).value)
            if ciudad_base:
                ciudad_base = ciudad_base.upper()

            transp_esp = clean_str(ws.cell(r, 6).value)
            zona_tecnica = clean_str(ws.cell(r, 7).value)
            new_so = clean_str(ws.cell(r, 8).value)
            jefe_zona = clean_str(ws.cell(r, 9).value)
            ing_soporte = clean_str(ws.cell(r, 10).value)
            estructura = clean_str(ws.cell(r, 11).value)
            alt_num, alt_txt = parse_float_safe(ws.cell(r, 12).value)
            facturadora = clean_str(ws.cell(r, 13).value)
            municipio = clean_str(ws.cell(r, 14).value)
            ubicacion = clean_str(ws.cell(r, 15).value)

            correo_so = clean_email(ws.cell(r, 16).value)
            correo_jz = clean_email(ws.cell(r, 17).value)
            correo_ing = clean_email(ws.cell(r, 18).value)

            if nombre in existing_sitios:
                # Actualizar sitio existente
                s = existing_sitios[nombre]
                s.zona = zona
                s.zona_tecnica = zona_tecnica
                s.ciudad_base = ciudad_base
                s.municipio = municipio
                s.ubicacion = ubicacion
                s.codigo_transporte_lpu = cod_transp
                s.km = km_num
                s.km_texto = km_txt
                s.transporte_especial = transp_esp
                s.estructura = estructura
                s.altura_estructura = alt_num
                s.altura_estructura_texto = alt_txt
                s.supervisor_operativo = new_so
                s.correo_so = correo_so
                s.jefe_zona = jefe_zona
                s.correo_jefe_zona = correo_jz
                s.ingeniero_soporte = ing_soporte
                s.correo_ing_soporte = correo_ing
                s.facturadora = facturadora
                updated_count += 1
            else:
                # Insertar nuevo sitio
                nuevo = Sitio(
                    nombre=nombre,
                    zona=zona,
                    zona_tecnica=zona_tecnica,
                    ciudad_base=ciudad_base,
                    municipio=municipio,
                    ubicacion=ubicacion,
                    codigo_transporte_lpu=cod_transp,
                    km=km_num,
                    km_texto=km_txt,
                    transporte_especial=transp_esp,
                    estructura=estructura,
                    altura_estructura=alt_num,
                    altura_estructura_texto=alt_txt,
                    supervisor_operativo=new_so,
                    correo_so=correo_so,
                    jefe_zona=jefe_zona,
                    correo_jefe_zona=correo_jz,
                    ingeniero_soporte=ing_soporte,
                    correo_ing_soporte=correo_ing,
                    facturadora=facturadora,
                    estado="activo"
                )
                db.add(nuevo)
                existing_sitios[nombre] = nuevo
                inserted_count += 1

            current_batch += 1
            if current_batch >= batch_size:
                db.commit()
                current_batch = 0

        # Commit de cambios restantes
        db.commit()
        print(f"  -> {inserted_count} sitios insertados nuevos.")
        print(f"  -> {updated_count} sitios actualizados.")
        print(f"  -> {skipped_count} filas vacías omitidas.")

        # 4. Vincular automáticamente OTs existentes con sus sitios correspondientes
        print("\n[Paso 4/4] Vinculando Órdenes de Trabajo (OTs) existentes con la base de datos de Sitios...")
        ots = db.query(Ot).all()
        vinculadas = 0
        for ot in ots:
            if not ot.sitio:
                continue
            # Buscar coincidencia exacta o por prefijo
            ot_sitio_upper = ot.sitio.upper().strip()
            # 1. Coincidencia exacta
            if ot_sitio_upper in existing_sitios:
                ot.sitio_id = existing_sitios[ot_sitio_upper].id
                vinculadas += 1
                continue
            # 2. Coincidencia por split (ej: 'ANT.APARTADO - EB Apartadó Centro (ANT-028)')
            parts = ot_sitio_upper.split(" - ")
            nombre_base = parts[0].strip()
            if nombre_base in existing_sitios:
                ot.sitio_id = existing_sitios[nombre_base].id
                vinculadas += 1
                continue

        db.commit()
        print(f"  -> {vinculadas} OTs vinculadas exitosamente con su Sitio maestro.")

        print("\n==================================================================")
        print(" [SUCCESS] ¡CARGA Y SINCRONIZACIÓN DE SITIOS FINALIZADA CON ÉXITO!")
        total_en_bd = db.query(Sitio).count()
        print(f" Total de Sitios registrados en 'sitios': {total_en_bd}")
        print("==================================================================")

    except Exception as e:
        db.rollback()
        print(f"\n[ERROR CRÍTICO] Ocurrió un error durante la carga: {e}")
        raise
    finally:
        db.close()

if __name__ == "__main__":
    load_sitios()
