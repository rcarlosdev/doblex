import openpyxl, glob, json, os

wb_tip = openpyxl.load_workbook(glob.glob('docs/*Tipolog*.xlsx')[0], data_only=True)
wb_fac = openpyxl.load_workbook('docs/ARCHIVO FACTURACIÓN.xlsx', data_only=True)
ws_fac = wb_fac['LISTA DE PRECIO 2026']

def extract_tip_sheet(ws):
    items = []
    for r in range(3, ws.max_row + 1):
        texto = ws.cell(r, 1).value
        codigo = ws.cell(r, 2).value
        alcance = ws.cell(r, 3).value
        cant = ws.cell(r, 4).value
        um = ws.cell(r, 5).value
        
        if not texto and not alcance:
            continue
            
        texto_str = str(texto or '').strip()
        alcance_str = str(alcance or '').strip()
        
        if 'Los items de instalaci' in texto_str:
            continue
            
        if 'CÓDIGO SAP' in texto_str.upper() or 'CODIGO SAP' in texto_str.upper():
            if '4X4' in texto_str.upper():
                codigo = 3042100
                texto_str = 'SRV TRANSPORTE VEHÍCULO 4X4 (LPU)'
                um = 'VIAJE'
            elif 'LANCHA' in texto_str.upper():
                codigo = 3042226
                texto_str = 'SRV TRANSPORTE FLUVIAL (LANCHA/PANGA/CANOA)'
                um = 'VIAJE'
            elif 'COTERO' in texto_str.upper() or 'SEMOVIENTE' in texto_str.upper():
                codigo = 3041820
                texto_str = 'SRV TRANSPORTE ESPECIAL (COTEROS O SEMOVIENTE)'
                um = 'DIA'

        tipo = 'Material'
        t_upper = texto_str.upper()
        if t_upper.startswith('MO ') or t_upper.startswith('SRV ') or 'CAMBIO ' in t_upper or 'INSTALACION' in t_upper or 'INSTAL' in t_upper or 'DESMONTE' in t_upper:
            tipo = 'Mano de Obra'
        elif t_upper.startswith('MT ') or t_upper.startswith('MF ') or 'SUMINISTRO' in t_upper or 'SUM ' in t_upper:
            tipo = 'Material'
            
        items.append({
            'codigo_sap': str(codigo).strip() if codigo else '',
            'texto_sap': texto_str,
            'alcance': alcance_str,
            'unidad': str(um or 'UNIDAD').strip(),
            'cantidad_sugerida': cant if (isinstance(cant, (int, float)) and cant > 0) else 1,
            'tipo': tipo
        })
    return items

tip1_items = extract_tip_sheet(wb_tip['TIPOLOGÍA 1'])
tip3_items = extract_tip_sheet(wb_tip['TIPOLOGÍA 3'])
tip5_items = extract_tip_sheet(wb_tip['TIPOLOGÍA 5'])

# Tipología 2 (PE) de LISTA DE PRECIO 2026
pe_items = []
pe_seen = set()
for r in range(3, ws_fac.max_row + 1):
    sap = ws_fac.cell(r, 3).value
    desc = str(ws_fac.cell(r, 4).value or '').strip()
    um = str(ws_fac.cell(r, 7).value or '').strip()
    cat = str(ws_fac.cell(r, 5).value or '').strip()
    if not sap or sap in pe_seen:
        continue
    d_upper = desc.upper()
    if any(k in d_upper for k in ['GE/ATS', 'PLANTA ELECT', 'GRUPO ELECTROGEN', 'FLEETGUARD', 'LF16015', 'FS1242', 'PRECALENTADOR', 'REGULADOR VOLTAJE', 'AVR SX460', 'BOMBA DE AGUA PERKINS', 'ACEITE LUBRICANTE 15W40', 'ANTICONGELANTE 50/50']):
        pe_seen.add(sap)
        tipo = 'Material' if ('MT ' in d_upper or 'SUM ' in d_upper or 'SUMINISTRO' in d_upper or cat == 'Material') else 'Mano de Obra'
        pe_items.append({
            'codigo_sap': str(sap),
            'texto_sap': desc,
            'alcance': desc,
            'unidad': um or 'UNIDAD',
            'cantidad_sugerida': 1,
            'tipo': tipo
        })

# Tipología 4 (PW) de LISTA DE PRECIO 2026
pw_items = []
pw_seen = set()
for r in range(3, ws_fac.max_row + 1):
    sap = ws_fac.cell(r, 3).value
    desc = str(ws_fac.cell(r, 4).value or '').strip()
    um = str(ws_fac.cell(r, 7).value or '').strip()
    cat = str(ws_fac.cell(r, 5).value or '').strip()
    if not sap or sap in pw_seen:
        continue
    d_upper = desc.upper()
    if any(k in d_upper for k in ['PW/BT', 'BATERIA', 'BANCO DE BATERIA', 'RECTIFICADOR', 'POWER CORE', 'GMT APLICA TODO POWER', 'FUSIBLE DC']):
        pw_seen.add(sap)
        tipo = 'Material' if ('MT ' in d_upper or 'SUM ' in d_upper or 'SUMINISTRO' in d_upper or cat == 'Material') else 'Mano de Obra'
        pw_items.append({
            'codigo_sap': str(sap),
            'texto_sap': desc,
            'alcance': desc,
            'unidad': um or 'UNIDAD',
            'cantidad_sugerida': 1,
            'tipo': tipo
        })

catalogo = {
    'TIPOLOGIA_1': {
        'id': 'TIPOLOGIA_1',
        'numero': 1,
        'nombre_corto': 'SPT',
        'nombre_completo': 'SPT - SISTEMA PUESTA A TIERRA',
        'subsistema_key': 'SPT - SISTEMA PUESTA A TIERRA',
        'items': tip1_items
    },
    'TIPOLOGIA_2': {
        'id': 'TIPOLOGIA_2',
        'numero': 2,
        'nombre_corto': 'PE',
        'nombre_completo': 'PE - GRUPO ELECTROGENO',
        'subsistema_key': 'PE - GRUPO ELECTROGENO',
        'items': pe_items
    },
    'TIPOLOGIA_3': {
        'id': 'TIPOLOGIA_3',
        'numero': 3,
        'nombre_corto': 'AA',
        'nombre_completo': 'AA - AIRES ACONDICIONADOS',
        'subsistema_key': 'AA - AIRES ACONDICIONADOS',
        'items': tip3_items
    },
    'TIPOLOGIA_4': {
        'id': 'TIPOLOGIA_4',
        'numero': 4,
        'nombre_corto': 'PW',
        'nombre_completo': 'PW - POWER',
        'subsistema_key': 'PW - POWER',
        'items': pw_items
    },
    'TIPOLOGIA_5': {
        'id': 'TIPOLOGIA_5',
        'numero': 5,
        'nombre_corto': 'MT-BT',
        'nombre_completo': 'MT-BT - MEDIA Y BAJA TENSION',
        'subsistema_key': 'MT-BT - MEDIA Y BAJA TENSION',
        'items': tip5_items
    }
}

os.makedirs('frontend/src/data', exist_ok=True)
out_path = 'frontend/src/data/catalogoTipologias.json'
with open(out_path, 'w', encoding='utf-8') as f:
    json.dump(catalogo, f, ensure_ascii=False, indent=2)

print(f'Guardado exitoso en {out_path}!')
for k, v in catalogo.items():
    print(f"{k}: {len(v['items'])} items")
